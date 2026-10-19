"""Read source Excel/CSV tables without modifying the originals."""
from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from .shared import clean, read_csv


@dataclass
class SourceTable:
    kind: str
    path: Path
    sheet: str
    rows: list[dict]


def table_kind(headers: set[str]) -> str:
    if {"Patient_No", "Order_Text", "Order_Creation_Date"} <= headers:
        return "medication"
    if {"Patient_No", "Adm_DateTime", "Dis_DateTime", "Case_No"} <= headers:
        return "visit"
    if {"Patient_No", "Patient_DOB", "Gender"} <= headers:
        return "person"
    if {"Patient_No", "Test_Code", "Test_Result"} <= headers:
        return "lab"
    if {"Patient_No", "Surgical_Code", "operation_date"} <= headers:
        return "procedure"
    return "unknown"


def discover(root: Path) -> list[SourceTable]:
    if not root.is_dir():
        raise ValueError(f"Input directory does not exist: {root}")
    tables: list[SourceTable] = []
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.name.startswith(("~$", ".")):
            continue
        if "__MACOSX" in path.parts:
            continue
        if path.suffix.lower() == ".csv":
            rows = read_csv(path)
            kind = table_kind(set(rows[0]) if rows else set())
            if kind != "unknown":
                for line, row in enumerate(rows, 2):
                    row["_source_row"] = line
                tables.append(SourceTable(kind, path.relative_to(root), "CSV", rows))
        elif path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in workbook:
                    iterator = sheet.iter_rows(values_only=True)
                    header = next(iterator, ())
                    named = [clean(value) for value in header]
                    duplicates = [h for h, n in Counter(named).items() if h and n > 1]
                    if duplicates:
                        raise ValueError(f"Duplicate headers in {path.name}/{sheet.title}: {duplicates}")
                    kind = table_kind(set(named))
                    if kind == "unknown":
                        continue
                    rows = []
                    for line, values in enumerate(iterator, 2):
                        row = {name: value for name, value in zip(named, values) if name}
                        # Excel formatting can extend to row 87 although only one is populated.
                        if not any(clean(value) for value in row.values()):
                            continue
                        row["_source_row"] = line
                        rows.append(row)
                    tables.append(SourceTable(kind, path.relative_to(root), sheet.title, rows))
            finally:
                workbook.close()
    return tables
